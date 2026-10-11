# Owner: Laptop B
"""Unit tests for Excel report generation (app/report/excel.py)."""

from __future__ import annotations

import tempfile
from pathlib import Path
from openpyxl import load_workbook

from app.report.excel import ExcelReport, append_excel_row
from app.schemas import (
    CaseRecord,
    Decision,
    DocumentRecord,
    DocumentType,
    FieldResult,
    FieldStatus,
    Issue,
    IssueKind,
    Route,
    WorkflowStatus,
)


def _sample_case(
    case_id: str = "case-12345678",
    route: Route = Route.PROCESS,
    sender_name: str = "Alice Smith",
    sender_email: str = "alice@example.com",
    extracted_name: str = "Alice Smith",
) -> CaseRecord:
    case = CaseRecord(
        case_id=case_id,
        scenario="job_application",
        sender_name=sender_name,
        sender_email=sender_email,
        workflow_status=WorkflowStatus.verified,
        route=route,
        decision=Decision(
            workflow_status=WorkflowStatus.verified,
            route=route,
            reasons=["All documents present and verified"],
            rule_ids=["R1"],
        ),
    )
    case.documents.append(
        DocumentRecord(
            doc_id="doc-1",
            filename="id_card.pdf",
            sha256="abc123sha",
            doc_type=DocumentType.id_document,
            pages=1,
            quality_verdict="ok",
            storage_path="data/originals/abc123sha.pdf",
        )
    )
    case.fields["id_document.applicant_name"] = FieldResult(
        field="applicant_name",
        value=extracted_name,
        status=FieldStatus.valid,
        confidence=0.95,
        ocr_value=extracted_name,
        vlm_value=extracted_name,
        page=0,
    )
    return case


def test_excel_report_sheets_and_atomic_write(tmp_path: Path):
    """Excel report creates 4 sheets, writes atomically, and copies to latest.xlsx."""
    run_id = "test-run-001"
    report = ExcelReport(run_id=run_id, data_dir=tmp_path)

    case1 = _sample_case(case_id="case-1", route=Route.PROCESS)
    case2 = _sample_case(case_id="case-2", route=Route.REQUEST_INFO, sender_name="Bob Jones", extracted_name="Robert Jones")
    case2.issues.append(
        Issue(
            field="degree_certificate",
            kind=IssueKind.missing_document,
            detail="Missing degree certificate",
            doc="degree_certificate",
        )
    )

    report.append_case(case1, email_status="dry_run", elapsed_s=1.5)
    report.append_case(case2, email_status="held", elapsed_s=2.5)

    report_path = report.finalise(elapsed_s=4.0, skipped_dups=1, dup_files=0)

    assert report_path.exists()
    latest_path = tmp_path / "reports" / "latest.xlsx"
    assert latest_path.exists()

    # Inspect sheets
    wb = load_workbook(str(report_path))
    assert "Applicants" in wb.sheetnames
    assert "Fields" in wb.sheetnames
    assert "Issues" in wb.sheetnames
    assert "Run" in wb.sheetnames

    # Applicants sheet check
    ws_app = wb["Applicants"]
    assert ws_app.max_row == 3  # Header + 2 cases

    headers = [cell.value for cell in ws_app[1]]
    assert "Case ID" in headers
    assert "Applicant Name" in headers
    assert "Name Flag" in headers
    assert "Route" in headers
    assert "Email Status" in headers

    row2 = [cell.value for cell in ws_app[2]]
    assert row2[0] == "case-1"
    assert row2[1] == "Alice Smith"
    assert not row2[2]  # No mismatch

    row3 = [cell.value for cell in ws_app[3]]
    assert row3[0] == "case-2"
    assert row3[2] == "MISMATCH"  # Bob Jones vs Robert Jones
    assert row3[7] == "REQUEST_INFO"
    assert row3[9] == "held"

    # Fields sheet check
    ws_fields = wb["Fields"]
    assert ws_fields.max_row >= 3
    f_headers = [cell.value for cell in ws_fields[1]]
    assert "Case ID" in f_headers
    assert "Field" in f_headers
    assert "Status" in f_headers

    # Issues sheet check
    ws_issues = wb["Issues"]
    assert ws_issues.max_row == 2  # Header + 1 issue from case 2
    iss_row = [cell.value for cell in ws_issues[2]]
    assert iss_row[0] == "case-2"
    assert iss_row[1] == "missing_document"
    assert iss_row[2] == "degree_certificate"

    # Run sheet check
    ws_run = wb["Run"]
    run_metrics = {cell[0].value: cell[1].value for cell in ws_run.iter_rows(min_row=2) if cell[0].value}
    assert run_metrics.get("Total Cases") == 2
    assert run_metrics.get("Duplicate Emails Skipped") == 1
    assert run_metrics.get("PROCESS") == 1
    assert run_metrics.get("REQUEST_INFO") == 1


def test_excel_route_colouring(tmp_path: Path):
    """Rows in Applicants sheet are coloured according to route."""
    report = ExcelReport(run_id="colors", data_dir=tmp_path)
    case_p = _sample_case("c-proc", Route.PROCESS)
    case_r = _sample_case("c-req", Route.REQUEST_INFO)
    case_m = _sample_case("c-man", Route.MANUAL_VERIFICATION)
    case_h = _sample_case("c-hum", Route.HUMAN_REVIEW)

    for c in [case_p, case_r, case_m, case_h]:
        report.append_case(c)

    path = report.finalise()
    wb = load_workbook(str(path))
    ws = wb["Applicants"]

    # Check that fill colors are assigned to data rows
    for row_idx in range(2, 6):
        fill = ws.cell(row=row_idx, column=1).fill
        assert fill.fill_type == "solid"
        assert fill.start_color.rgb is not None


def test_legacy_append_excel_row(tmp_path: Path):
    """Legacy helper append_excel_row functions properly."""
    custom_path = tmp_path / "reports" / "custom.xlsx"
    case = _sample_case("leg-1", Route.PROCESS)
    append_excel_row(case, report_path=custom_path)
    assert custom_path.exists()
